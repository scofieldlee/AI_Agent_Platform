"""文档处理器 — PDF / TXT / Markdown / Word 等文本类素材（Phase 7）.

- PDF: pypdf 逐页提取文本
- TXT / Markdown / CSV: 直接读取
- Word (docx): python-docx
- 长文本按字符分块，每块一个 KnowledgeUnit（unit_type=text）
- 调用 Model Center LLM 生成整体摘要 + 标签（失败时降级为统计描述，不阻断流程）
- 文本块由索引服务走文本 Embedding（本地 Qwen3-VL-Embedding / DashScope 均支持）
"""

import logging
import os
from typing import Dict, Any, List, Optional

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.multimodal.constants import (AssetStatus, UnitType, UnitStatus,
                                      TagSource)
from app.multimodal.repositories import asset_repo, knowledge_unit_repo
from app.multimodal.services import storage_service

logger = logging.getLogger(__name__)

# 单个知识单元的字符数（分块目标大小 / 硬上限）
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 120
MAX_CHARS_PER_UNIT = 4000   # Embedding 输入安全上限
MAX_UNITS_PER_ASSET = 200   # 防止超长文档撑爆任务


# ============================================================
# 文本提取
# ============================================================

def extract_pdf_pages(abs_path: str) -> List[Dict[str, Any]]:
    """pypdf 逐页提取。返回 [{page, text}]（空页跳过）。"""
    from pypdf import PdfReader
    reader = PdfReader(abs_path)
    pages: List[Dict[str, Any]] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append({"page": i + 1, "text": text})
    return pages


def extract_plain_text(abs_path: str, max_bytes: int = 8 * 1024 * 1024) -> str:
    """txt / md / csv 直接读取。"""
    size = os.path.getsize(abs_path)
    if size > max_bytes:
        raise ValueError(f"文本文件过大: {size} bytes")
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            with open(abs_path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("无法识别的文本编码")


def extract_docx(abs_path: str) -> str:
    """python-docx 提取段落文本。"""
    from docx import Document
    doc = Document(abs_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 表格内容也提取
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_document_text(abs_path: str, file_type: str) -> Dict[str, Any]:
    """按类型提取文本。返回 {pages: [{page, text}] | None, text: str}。"""
    lower = abs_path.lower()
    if file_type == "pdf" or lower.endswith(".pdf"):
        pages = extract_pdf_pages(abs_path)
        return {"pages": pages, "text": "\n\n".join(p["text"] for p in pages)}
    if lower.endswith(".docx"):
        return {"pages": None, "text": extract_docx(abs_path)}
    # txt / md / csv / 其它文本
    return {"pages": None, "text": extract_plain_text(abs_path)}


def chunk_text(text: str, size: int = CHUNK_SIZE,
               overlap: int = CHUNK_OVERLAP) -> List[str]:
    """按字符滑动窗口分块。"""
    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks: List[str] = []
    step = max(size - overlap, 1)
    start = 0
    while start < len(text) and len(chunks) < MAX_UNITS_PER_ASSET:
        chunks.append(text[start:start + size])
        start += step
    return chunks


# ============================================================
# AI 摘要（可选增强，失败不阻断）
# ============================================================

SUMMARY_SYSTEM_PROMPT = (
    "你是文档分析助手。根据用户给出的文档内容，输出严格的 JSON：\n"
    '{"description": "一句话概括（50字内）", "summary": "内容摘要（200字内）", '
    '"tags": ["标签1", "标签2", "标签3"]}\n'
    "标签用中文，3-6 个，体现文档主题/领域/类型。只输出 JSON，不要多余文字。"
)


async def summarize_document(filename: str, text: str) -> Optional[Dict[str, Any]]:
    """调用 Model Center 生成文档摘要与标签。失败返回 None。"""
    if not text:
        return None
    sample = text[:6000]
    try:
        from app.models_center.service import ModelService
        service = ModelService()
        result = await service.chat(
            system_prompt=SUMMARY_SYSTEM_PROMPT,
            user_prompt=f"文件名: {filename}\n\n文档内容:\n{sample}",
            temperature=0.2,
            max_tokens=800,
        )
        content = (result or {}).get("content") or ""
        return _parse_summary_json(content)
    except Exception as e:
        logger.warning(f"Document LLM summary failed (degrade to stats): {e}")
        return None


def _parse_summary_json(content: str) -> Optional[Dict[str, Any]]:
    """从 LLM 输出中解析 JSON（容忍 ```json 包裹 / 前后缀文字）。"""
    import json
    import re
    if not content:
        return None
    m = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", content, re.S)
    raw = m.group(1) if m else content
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        data = json.loads(raw[start:end + 1])
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    tags = data.get("tags")
    if not isinstance(tags, list):
        tags = []
    return {
        "description": str(data.get("description") or "")[:200],
        "summary": str(data.get("summary") or "")[:1000],
        "tags": [str(t).strip() for t in tags if str(t).strip()][:6],
    }


# ============================================================
# 主流程
# ============================================================

async def process(db: AsyncSession, asset_id: int) -> Dict[str, Any]:
    """文档素材处理主流程: 文本提取 → 分块 → units →（可选）LLM 摘要 → REVIEW_REQUIRED."""
    asset = await asset_repo.get_asset(db, asset_id)
    if not asset:
        return {"success": False, "error": f"Asset {asset_id} not found"}
    if not asset.storage_path:
        return {"success": False, "error": "Asset has no storage_path"}

    abs_path = storage_service.abs_file_path(asset.storage_path)
    kb_id, aid = asset.knowledge_base_id, asset.id
    filename = asset.original_filename or asset.name or ""

    # 1. 文本提取
    try:
        extracted = extract_document_text(abs_path, asset.file_type or "")
    except Exception as e:
        return {"success": False, "error": f"文档解析失败: {e}"}

    full_text = (extracted.get("text") or "").strip()
    if not full_text:
        return {"success": False,
                "error": "文档无可提取文本（可能是扫描件/图片型 PDF，需要 OCR 能力）"}

    pages: Optional[List[Dict[str, Any]]] = extracted.get("pages")
    system_meta = (await asset_repo.get_all_metadata(db, asset_id)).get("system", {})
    system_meta.update({
        "text_chars": len(full_text),
        "page_count": len(pages) if pages else None,
        "chunk_size": CHUNK_SIZE,
    })
    system_meta = {k: v for k, v in system_meta.items() if v is not None}
    await asset_repo.upsert_metadata(db, asset_id, "system", system_meta)

    # 2. 分块 → 知识单元
    await knowledge_unit_repo.delete_units_by_asset(db, asset_id)
    if pages:
        # PDF：按页分块，页内过长再切
        chunks: List[Dict[str, Any]] = []
        for p in pages:
            page_no, page_text = p["page"], p["text"]
            pieces = chunk_text(page_text) or [""]
            for j, piece in enumerate(pieces):
                chunks.append({"text": piece, "page": page_no,
                               "label": f"第 {page_no} 页" + (f"（{j + 1}）" if len(pieces) > 1 else "")})
    else:
        pieces = chunk_text(full_text)
        chunks = [{"text": t, "page": None, "label": f"片段 {i + 1}"}
                  for i, t in enumerate(pieces)]
    chunks = chunks[:MAX_UNITS_PER_ASSET]

    for i, chunk in enumerate(chunks):
        content = chunk["text"][:MAX_CHARS_PER_UNIT]
        description = (f"{filename} · {chunk['label']}"
                       if chunk["text"] else None)
        unit = await knowledge_unit_repo.create_unit(
            db, knowledge_base_id=kb_id, asset_id=aid,
            unit_type=UnitType.TEXT, unit_index=i,
            content=content, description=description,
            meta={"page": chunk["page"], "label": chunk["label"]},
            status=UnitStatus.PENDING)

    # 3. AI 摘要（可选增强，失败降级为统计描述）
    summary = await summarize_document(filename, full_text)
    if summary and summary.get("description"):
        ai_meta = {
            "description": summary["description"],
            "summary": summary.get("summary", ""),
            "tags": summary.get("tags", []),
            "analyzed_by": "llm",
        }
    else:
        ai_meta = {
            "description": (
                f"{'PDF' if asset.file_type == 'pdf' else '文档'}素材，"
                f"共 {len(pages) if pages else '-'} 页 / {len(full_text)} 字，"
                f"已切分为 {len(chunks)} 个检索单元"),
            "tags": ["文档", "PDF" if asset.file_type == "pdf" else "文本"],
            "analyzed_by": "stats",
        }
    await asset_repo.upsert_metadata(db, asset_id, "ai", ai_meta)
    ai_tags = [t for t in ai_meta.get("tags", []) if t]
    if ai_tags:
        await asset_repo.replace_ai_tags(db, asset_id, ai_tags)

    # 4. 状态 → REVIEW_REQUIRED（审核后进入索引）
    await asset_repo.update_asset(db, asset, {"status": AssetStatus.REVIEW_REQUIRED})
    return {"success": True, "units": len(chunks),
            "text_chars": len(full_text),
            "pages": len(pages) if pages else None,
            "summary_source": ai_meta.get("analyzed_by")}
